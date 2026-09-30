import io
from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity

from server.models import Document

export_bp = Blueprint("export", __name__, url_prefix="/api/documents")


def _get_current_user_id() -> int:
    identity = get_jwt_identity()
    try:
        return int(identity)
    except (TypeError, ValueError):
        return identity


@export_bp.get("/<int:doc_id>/export")
@jwt_required()
def export_document(doc_id):
    """Query param: ?format=pdf|docx|txt (default docx)"""
    user_id = _get_current_user_id()
    doc = Document.query.filter_by(id=doc_id, user_id=user_id).first()
    if not doc:
        return jsonify({"error": "Document not found."}), 404

    fmt = (request.args.get("format") or "docx").lower()
    filename_base = "".join(
        c for c in (doc.title or "document") if c.isalnum() or c in (" ", "-", "_")
    ).strip().replace(" ", "_") or "document"

    try:
        if fmt == "docx":
            buffer = _build_docx(doc)
            return send_file(
                buffer,
                as_attachment=True,
                download_name=f"{filename_base}.docx",
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        elif fmt == "pdf":
            buffer = _build_pdf(doc)
            return send_file(
                buffer,
                as_attachment=True,
                download_name=f"{filename_base}.pdf",
                mimetype="application/pdf",
            )
        elif fmt == "txt":
            buffer = io.BytesIO(doc.content.encode("utf-8"))
            return send_file(
                buffer,
                as_attachment=True,
                download_name=f"{filename_base}.txt",
                mimetype="text/plain",
            )
        else:
            return jsonify({"error": "format must be 'pdf', 'docx', or 'txt'."}), 400
    except Exception as e:
        return jsonify({"error": f"Export generation failed: {str(e)}"}), 500


def _build_docx(doc) -> io.BytesIO:
    try:
        from docx import Document as DocxDocument
        docx_file = DocxDocument()
        docx_file.add_heading(doc.title or "Legal Document", level=1)

        for line in (doc.content or "").split("\n"):
            stripped = line.strip()
            if stripped:
                docx_file.add_paragraph(stripped)
            else:
                docx_file.add_paragraph("")

        buffer = io.BytesIO()
        docx_file.save(buffer)
        buffer.seek(0)
        return buffer
    except ImportError:
        # Fallback if python-docx not installed
        buffer = io.BytesIO(f"# {doc.title}\n\n{doc.content}".encode("utf-8"))
        buffer.seek(0)
        return buffer


def _build_pdf(doc) -> io.BytesIO:
    try:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch

        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(
            buffer,
            pagesize=LETTER,
            topMargin=inch,
            bottomMargin=inch,
            leftMargin=inch,
            rightMargin=inch,
        )
        styles = getSampleStyleSheet()
        story = [Paragraph(doc.title or "Legal Document", styles["Title"]), Spacer(1, 18)]

        for line in (doc.content or "").split("\n"):
            stripped = line.strip()
            if stripped:
                safe = (
                    stripped.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )
                story.append(Paragraph(safe, styles["BodyText"]))
                story.append(Spacer(1, 8))

        pdf.build(story)
        buffer.seek(0)
        return buffer
    except ImportError:
        # Fallback if reportlab not installed
        buffer = io.BytesIO(f"%PDF-1.4\n% Plaintext fallback for {doc.title}\n{doc.content}".encode("utf-8"))
        buffer.seek(0)
        return buffer
